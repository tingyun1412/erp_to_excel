"""
廠商帳號 / 發票跳過名單 / 發票開立記錄 / 出貨提醒：本機檔案儲存。
取代原本的 Google Sheets，不再需要任何 Google 憑證。

發票開立記錄需要「先查是否已開過、再送出、再寫入」全程不能有兩人同時通過檢查，
用 filelock 包住「認領」這個步驟（見 try_claim_order），避免同一張銷貨單被
兩人同時送出、開出兩張真實發票。其餘幾張表（廠商帳號/跳過名單）很少同時被
兩人編輯，直接讀寫本機 JSON 檔即可，沒有另外加鎖。

出貨提醒不再自己存一份，直接即時讀取倉管本來就在維護的那份共用 Excel
（Z:\\倉管\\訂購單預交-每週出貨.xlsx 的「出貨要求」工作表），不需要另外匯入。
"""
import json
from datetime import datetime, timezone, timedelta
from pathlib import Path

import streamlit as st
from filelock import FileLock, Timeout

_TW = timezone(timedelta(hours=8))
_BASE_DIR = Path(__file__).resolve().parent / "local_data"

_VENDORS_FILE = _BASE_DIR / "vendors.json"
_INVOICE_SKIP_FILE = _BASE_DIR / "invoice_skip.json"
_EINVOICE_LOG_FILE = _BASE_DIR / "einvoice_log.json"
_EINVOICE_LOG_LOCK = _BASE_DIR / "einvoice_log.lock"

# 出貨提醒直接讀這份共用檔案，不再自己存一份／提供匯入功能
SHIPPING_NOTES_PATH = Path(r"Z:\倉管\訂購單預交-每週出貨.xlsx")
_SHIPPING_NOTES_SHEET = "出貨要求"

# 「處理中」狀態超過這個時間還沒轉成「已開立」或「失敗」，視為呼叫端已經當掉/中斷，
# 允許重新認領，避免卡死的紀錄永遠擋住這張銷貨單無法重新送出。
_CLAIM_STALE_MINUTES = 10


def _load(path: Path) -> list[dict]:
    if not path.exists():
        return []
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return []


def _save(path: Path, records: list[dict]):
    _BASE_DIR.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")


# ── 廠商帳號 ──────────────────────────────────────────────────────

def load_vendors() -> list[dict]:
    return _load(_VENDORS_FILE)


def save_vendor(name: str, url: str, username: str, password: str):
    records = _load(_VENDORS_FILE)
    for r in records:
        if r.get("公司名稱") == name:
            r.update({"網址": url, "帳號": username, "密碼": password})
            _save(_VENDORS_FILE, records)
            return
    records.append({"公司名稱": name, "網址": url, "帳號": username, "密碼": password})
    _save(_VENDORS_FILE, records)


def delete_vendor(name: str):
    records = [r for r in _load(_VENDORS_FILE) if r.get("公司名稱") != name]
    _save(_VENDORS_FILE, records)


# ── 電子發票：跳過名單 ──────────────────────────────────────────────

def load_invoice_skip_list() -> list[dict]:
    return _load(_INVOICE_SKIP_FILE)


def save_invoice_skip(customer_name: str, reason: str):
    records = _load(_INVOICE_SKIP_FILE)
    for r in records:
        if r.get("客戶名稱") == customer_name:
            r["原因"] = reason
            _save(_INVOICE_SKIP_FILE, records)
            return
    records.append({"客戶名稱": customer_name, "原因": reason})
    _save(_INVOICE_SKIP_FILE, records)


def delete_invoice_skip(customer_name: str):
    records = [r for r in _load(_INVOICE_SKIP_FILE) if r.get("客戶名稱") != customer_name]
    _save(_INVOICE_SKIP_FILE, records)


# ── 電子發票：逐張開立記錄 ────────────────────────────────────────────

def load_einvoice_log() -> list[dict]:
    return _load(_EINVOICE_LOG_FILE)


def _is_stale_claim(row: dict) -> bool:
    if row.get("狀態") != "處理中":
        return False
    try:
        claimed_at = datetime.strptime(row["開立時間"], "%Y/%m/%d %H:%M").replace(tzinfo=_TW)
    except Exception:
        return True
    return (datetime.now(_TW) - claimed_at) > timedelta(minutes=_CLAIM_STALE_MINUTES)


def try_claim_order(order_no: str, customer_name: str, tax_id: str) -> bool:
    """
    原子性地嘗試「認領」這張銷貨單準備送出發票。
    成功回傳 True（呼叫端接著可以去網站送出）；已被認領中或已開立成功回傳 False
    （呼叫端應該跳過，不要重複送出）。認領成功後，呼叫端無論送出成功或失敗，
    都必須呼叫 save_einvoice_log 寫入最終狀態，否則這張單會卡在「處理中」
    （最多卡 _CLAIM_STALE_MINUTES 分鐘後可重新認領）。
    """
    _BASE_DIR.mkdir(parents=True, exist_ok=True)
    try:
        with FileLock(str(_EINVOICE_LOG_LOCK), timeout=30):
            records = _load(_EINVOICE_LOG_FILE)
            for r in records:
                if r.get("銷貨單號") == order_no:
                    if r.get("狀態") == "已開立":
                        return False
                    if r.get("狀態") == "處理中" and not _is_stale_claim(r):
                        return False
                    break
            now = datetime.now(_TW).strftime("%Y/%m/%d %H:%M")
            row = {"銷貨單號": order_no, "客戶名稱": customer_name, "統一編號": tax_id,
                   "發票號碼": "", "狀態": "處理中", "開立時間": now, "備註": ""}
            for i, r in enumerate(records):
                if r.get("銷貨單號") == order_no:
                    records[i] = row
                    _save(_EINVOICE_LOG_FILE, records)
                    return True
            records.append(row)
            _save(_EINVOICE_LOG_FILE, records)
            return True
    except Timeout:
        # 鎖不到代表有其他人正在動這份記錄，保守起見當作認領失敗，不要搶著送出
        return False


def save_einvoice_log(order_no: str, customer_name: str, tax_id: str,
                       invoice_no: str, status: str, remark: str = ""):
    """依銷貨單號 upsert 最終狀態。狀態："已開立"（成功）／"失敗"（可重試）。"""
    _BASE_DIR.mkdir(parents=True, exist_ok=True)
    with FileLock(str(_EINVOICE_LOG_LOCK), timeout=30):
        records = _load(_EINVOICE_LOG_FILE)
        now = datetime.now(_TW).strftime("%Y/%m/%d %H:%M")
        row = {"銷貨單號": order_no, "客戶名稱": customer_name, "統一編號": tax_id,
               "發票號碼": invoice_no or "", "狀態": status, "開立時間": now, "備註": remark}
        for i, r in enumerate(records):
            if r.get("銷貨單號") == order_no:
                records[i] = row
                _save(_EINVOICE_LOG_FILE, records)
                return
        records.append(row)
        _save(_EINVOICE_LOG_FILE, records)


def delete_einvoice_log(order_no: str):
    with FileLock(str(_EINVOICE_LOG_LOCK), timeout=30):
        records = [r for r in _load(_EINVOICE_LOG_FILE) if r.get("銷貨單號") != order_no]
        _save(_EINVOICE_LOG_FILE, records)


# ── 出貨提醒（直接讀倉管維護的共用檔案，不寫入、不提供編輯）───────────────

@st.cache_data(ttl=60)
def load_shipping_notes() -> list[dict]:
    """
    即時讀取 Z:\\倉管\\訂購單預交-每週出貨.xlsx 的「出貨要求」工作表。
    這是網路磁碟上的共用檔案，讀取有實際延遲，所以短暫快取 60 秒，
    避免每次網頁互動都重新讀一次網路磁碟。
    讀不到（檔案不存在／被鎖住／沒有權限／格式不對）時丟例外，由呼叫端決定要不要顯示。
    """
    import openpyxl
    if not SHIPPING_NOTES_PATH.exists():
        raise FileNotFoundError(f"找不到出貨提醒檔案：{SHIPPING_NOTES_PATH}")
    wb = openpyxl.load_workbook(SHIPPING_NOTES_PATH, data_only=True, read_only=True)
    if _SHIPPING_NOTES_SHEET not in wb.sheetnames:
        raise ValueError(f"{SHIPPING_NOTES_PATH.name} 裡找不到「{_SHIPPING_NOTES_SHEET}」工作表")
    ws = wb[_SHIPPING_NOTES_SHEET]
    rows = []
    for r in range(2, ws.max_row + 1):
        cust = ws.cell(r, 1).value
        if not cust:
            continue
        rows.append({
            "客戶": str(cust).strip(),
            "出貨要求": str(ws.cell(r, 2).value or "").strip(),
            "備註": str(ws.cell(r, 3).value or "").strip(),
        })
    return rows
